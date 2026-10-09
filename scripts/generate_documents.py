from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = ROOT.parents[1] / "outputs"
DELIVERY = ROOT / "docs" / "entrega"
DIAGRAMS = ROOT / "docs" / "diagrams"
SHOTS = ROOT / "docs" / "screenshots"
TEMPLATE = Path(
    r"C:\Users\usuário\OneDrive\Documentos\FACULDADE\7 SEMESTRE\Projeto Integrador Transdisciplinar em Engenharia de Software I I\PIT_atividade.docx"
)
REPO_URL = "https://github.com/Mathauz/doce-encanto"

PINK = "B73865"
DARK = "3A2428"
CREAM = "FFF8F2"
ROSE = colors.HexColor("#B73865")
COCOA = colors.HexColor("#3A2428")
PALE = colors.HexColor("#FFF1F5")


stories = [
    ("US01", "Criar conta", "visitante", "criar uma conta com usuário e senha", "acompanhar meus pedidos", "Dados obrigatórios válidos; usuário não pode se repetir; senha confirmada.", "Alta", 3),
    ("US02", "Entrar no sistema", "cliente", "autenticar-me", "acessar checkout e histórico", "Credenciais válidas redirecionam; inválidas exibem erro sem revelar qual campo falhou.", "Alta", 2),
    ("US03", "Sair do sistema", "cliente", "encerrar minha sessão", "proteger minha conta", "Logout remove autenticação e mantém navegação pública disponível.", "Alta", 1),
    ("US04", "Ver catálogo", "visitante", "consultar cupcakes ativos", "escolher o que comprar", "Somente ativos aparecem; nome, descrição, imagem e preço são exibidos.", "Alta", 3),
    ("US05", "Adicionar ao carrinho", "visitante", "adicionar um cupcake", "montar meu pedido", "Produto ativo é incluído; quantidade final fica entre 1 e 20.", "Alta", 3),
    ("US06", "Alterar quantidade", "visitante", "ajustar a quantidade de um item", "corrigir meu pedido", "Quantidade aceita de 1 a 20; total é recalculado no servidor.", "Alta", 2),
    ("US07", "Remover item", "visitante", "remover um item do carrinho", "desistir de um produto", "Item deixa a sessão e totais são atualizados.", "Alta", 1),
    ("US08", "Revisar carrinho", "visitante", "ver itens, quantidades e total", "confirmar a compra", "Carrinho vazio tem estado orientativo; preços vêm do banco.", "Alta", 3),
    ("US09", "Finalizar pedido", "cliente", "informar endereço e pagamento simulado", "registrar meu pedido", "Checkout autenticado, validado e transacional cria pedido e itens completos.", "Alta", 8),
    ("US10", "Informar observações", "cliente", "registrar instruções de entrega", "personalizar orientações", "Campo opcional aceita texto limitado e é preservado no pedido.", "Média", 1),
    ("US11", "Ver meus pedidos", "cliente", "consultar meu histórico", "acompanhar compras", "Lista contém apenas pedidos do usuário autenticado.", "Alta", 3),
    ("US12", "Ver detalhe do pedido", "cliente", "abrir um pedido meu", "conferir itens e status", "Acesso a pedido alheio retorna 404; snapshot de preços é exibido.", "Alta", 3),
    ("US13", "Gerenciar produtos", "administrador", "cadastrar, editar e inativar cupcakes", "manter o catálogo", "Admin permite CRUD; produto inativo não é vendido.", "Alta", 5),
    ("US14", "Atualizar status", "administrador", "alterar o andamento do pedido", "informar a operação", "Status limitado a recebido, preparo, entrega, concluído ou cancelado.", "Alta", 3),
    ("US15", "Usar em diferentes telas", "usuário", "navegar pelo celular ou computador", "comprar com conforto", "Layout responsivo sem rolagem horizontal em 390 px e funcional em desktop.", "Média", 5),
]

tasks = [
    ("Fundação", "Projeto Django, configurações, modelos e migrações", 8),
    ("Contas", "Cadastro, login, logout e permissões", 5),
    ("Catálogo", "Listagem, destaques, imagens e administração", 6),
    ("Carrinho", "Sessão, adicionar, atualizar, remover e totais", 8),
    ("Pedidos", "Checkout transacional, histórico e detalhes", 10),
    ("Qualidade", "Testes, cobertura, smoke responsivo e correções", 9),
    ("Entrega", "README, UML, manual, Render e roteiros", 10),
]

qa = [
    ("QA-01 — sessão interna simulada", "09/10/2026", "Cadastro, login, catálogo e saída funcionaram de ponta a ponta.", "Nenhuma falha impeditiva. A ordem dos destaques não seguia a curadoria visual; corrigida para preservar a ordem de cadastro.", "Pagamento real, fora do escopo."),
    ("QA-02 — sessão interna simulada", "09/10/2026", "Carrinho incluiu, atualizou e removeu itens; total permaneceu consistente.", "Limite acima de 20 foi bloqueado/corrigido no servidor e coberto por teste.", "Cupons, fora do escopo."),
    ("QA-03 — sessão interna simulada", "09/10/2026", "Checkout criou pedido, esvaziou o carrinho e preservou preço/descrição.", "Formulário inválido mantém os dados e explica os campos; sem defeito remanescente.", "Gateway de pagamento, explicitamente simulado."),
    ("QA-04 — sessão interna simulada", "09/10/2026", "Histórico e detalhe exibiram somente pedidos do usuário conectado.", "Tentativa de abrir pedido de outro usuário retorna 404, validada por teste automatizado.", "Rastreamento por mapa, fora do escopo."),
    ("QA-05 — sessão interna simulada", "09/10/2026", "Homepage e catálogo aprovados em 1440×1000 e 390×844.", "Foi encontrado erro 404 de favicon no smoke inicial; criado favicon local e console ficou sem erros.", "Aplicativo nativo, fora do escopo."),
]


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold=False, color=DARK, size=9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_doc_defaults(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for name, size, color in (("Title", 25, PINK), ("Heading 1", 17, PINK), ("Heading 2", 13, DARK), ("Heading 3", 11, PINK)):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)


def add_header_footer(doc: Document) -> None:
    for sec in doc.sections:
        hp = sec.header.paragraphs[0]
        hp.text = "DOCE ENCANTO  •  PROJETO INTEGRADOR TRANSDISCIPLINAR II"
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in hp.runs:
            r.font.name = "Arial"; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(PINK)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.add_run("Matheus da Silva Mattos  •  RGM 36338770  •  São Paulo, 2026")
        for r in fp.runs:
            r.font.name = "Arial"; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(DARK)


def add_title(doc: Document, title: str, subtitle: str = "") -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(140)
    r = p.add_run(title)
    r.bold = True; r.font.name = "Arial"; r.font.size = Pt(28); r.font.color.rgb = RGBColor.from_string(PINK)
    if subtitle:
        p2 = doc.add_paragraph(subtitle)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r2 in p2.runs:
            r2.font.name = "Arial"; r2.font.size = Pt(14); r2.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph("\nMatheus da Silva Mattos\nRGM 36338770\n\nSão Paulo\n2026").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_heading(doc: Document, text: str, level=1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15


def add_bullets(doc: Document, lines) -> None:
    for line in lines:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(line)


def add_table(doc: Document, headers, rows, widths=None, font_size=8.5) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", size=font_size)
        shade(table.rows[0].cells[i], PINK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            if len(table.rows) % 2 == 0:
                shade(cells[i], CREAM)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()


def add_picture(doc: Document, path: Path, width=6.2, caption="") -> None:
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        if caption:
            cp = doc.add_paragraph(caption)
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in cp.runs:
                r.italic = True; r.font.size = Pt(8)


def build_main_docx(path: Path) -> None:
    doc = Document()
    set_doc_defaults(doc); add_header_footer(doc)
    add_title(doc, "DOCE ENCANTO", "E-commerce acadêmico de cupcakes — documentação de projeto e implementação")

    add_heading(doc, "Resumo executivo")
    add_body(doc, "O Doce Encanto é uma aplicação web responsiva para venda demonstrativa de cupcakes. A solução permite cadastro e autenticação, consulta ao catálogo, carrinho em sessão, checkout com pagamento simulado, histórico de pedidos e administração do catálogo e do fluxo operacional. Foi desenvolvida em Django 5.2 LTS, com SQLite no ambiente local e configuração para PostgreSQL no Render.")
    add_body(doc, "Esta documentação consolida os artefatos das etapas I e II: requisitos, histórias, backlog, mapas, UML, banco de dados, interface, testes, laudo de qualidade, implantação, manual e roteiros de apresentação. O projeto é acadêmico; não processa pagamentos reais nem deve ser tratado como comércio pronto para produção.")
    add_heading(doc, "1. Visão, objetivos e escopo")
    add_body(doc, "Objetivo geral: construir um fluxo de compra simples, confiável e visualmente coerente com uma confeitaria artesanal, demonstrando competências de análise, arquitetura, persistência, implementação, testes e publicação.")
    add_bullets(doc, [
        "Público-alvo: consumidores que desejam encomendar cupcakes de forma rápida.",
        "Valor entregue: catálogo claro, carrinho persistente durante a sessão e acompanhamento do pedido.",
        "Incluído: contas, catálogo, carrinho, checkout simulado, pedidos e administração nativa do Django.",
        "Fora do escopo: gateway real, cupons, avaliações públicas, frete por distância, estoque complexo e aplicativo nativo.",
    ])
    add_heading(doc, "2. Metodologia e organização")
    add_body(doc, "O trabalho utilizou backlog priorizado, estimativas em pontos, implementação incremental, testes automatizados e sessões internas de QA simuladas. O monólito Django separa responsabilidades por modelos, formulários, serviços de carrinho, views, templates e administração.")
    add_table(doc, ["Etapa", "Atividade", "Estimativa (h)"], tasks, [3.2, 10.5, 2.8])
    add_picture(doc, DIAGRAMS / "mapa-afinidade.png", caption="Figura 1 — Mapa de afinidade dos requisitos.")

    add_heading(doc, "3. Histórias de usuário e critérios de aceitação")
    add_body(doc, "As histórias abaixo foram refinadas com prioridade e pontos relativos. As regras transversais são: preço sempre recalculado no servidor; quantidade de 1 a 20; produto inativo não pode ser comprado; pedido pertence a um único cliente; checkout é atômico; pagamento é demonstrativo.")
    for sid, title, actor, action, benefit, acceptance, priority, points in stories:
        add_heading(doc, f"{sid} — {title}", level=2)
        add_body(doc, f"Como {actor}, quero {action}, para {benefit}.")
        add_table(doc, ["Critérios de aceitação", "Prioridade", "Pontos"], [(acceptance, priority, points)], [12.3, 2.4, 1.8], 8)
    add_heading(doc, "4. Backlog priorizado")
    add_table(doc, ["ID", "História", "Prioridade", "Pontos"], [(s[0], s[1], s[6], s[7]) for s in stories], [2, 9.5, 2.6, 1.8], 8)

    add_heading(doc, "5. Requisitos")
    add_heading(doc, "5.1 Requisitos funcionais", level=2)
    add_bullets(doc, [
        "RF01–RF03: cadastrar, autenticar e encerrar sessão de usuários.",
        "RF04–RF08: exibir catálogo e operar carrinho em sessão.",
        "RF09–RF12: validar checkout, registrar pedido e restringir histórico/detalhe ao proprietário.",
        "RF13–RF14: administrar produtos e status de pedidos pelo Django Admin.",
        "RF15: oferecer páginas 404 e estados vazios compreensíveis.",
    ])
    add_heading(doc, "5.2 Requisitos não funcionais", level=2)
    add_bullets(doc, [
        "RNF01: interface responsiva a partir de 390 px, com contraste e foco visível.",
        "RNF02: segredos e configuração de produção fornecidos por variáveis de ambiente.",
        "RNF03: banco PostgreSQL em produção e SQLite apenas no desenvolvimento.",
        "RNF04: respostas sensíveis protegidas por autenticação e isolamento por usuário.",
        "RNF05: código reproduzível, com README, dependências fixadas, seed e CI.",
        "RNF06: testes automatizados no GitHub Actions e cobertura registrada.",
    ])
    add_heading(doc, "5.3 Regras de negócio", level=2)
    add_table(doc, ["Regra", "Descrição"], [
        ("RN01", "Quantidade por item entre 1 e 20."),
        ("RN02", "Somente produtos ativos podem entrar no carrinho/checkout."),
        ("RN03", "Preço e total são recalculados a partir do banco no servidor."),
        ("RN04", "Pedido guarda descrição e preço unitário no momento da compra."),
        ("RN05", "Checkout cria pedido e itens em uma transação única."),
        ("RN06", "Formas: Pix simulado, cartão na entrega e dinheiro."),
        ("RN07", "Status: recebido, em preparo, saiu para entrega, concluído ou cancelado."),
        ("RN08", "Cliente não visualiza pedido pertencente a outro usuário."),
    ], [2.4, 14.8])

    add_heading(doc, "6. Modelagem UML")
    add_picture(doc, DIAGRAMS / "casos-de-uso.png", caption="Figura 2 — Diagrama de casos de uso.")
    add_heading(doc, "6.1 Casos de uso expandidos", level=2)
    expanded = [
        ("UC01 — Cadastrar cliente", "Visitante", "Não estar autenticado", "Informa usuário, e-mail e senha; sistema valida, cria conta, autentica e redireciona.", "Conta criada e sessão iniciada.", "Usuário duplicado ou senha inválida mantém formulário com mensagens."),
        ("UC02 — Finalizar pedido", "Cliente", "Estar autenticado e possuir carrinho válido", "Revisa itens; informa endereço, CEP, pagamento e observações; sistema revalida catálogo e totais; cria pedido e itens atomicamente; esvazia carrinho.", "Pedido com status Recebido disponível no histórico.", "Carrinho vazio, produto inativo ou formulário inválido impedem criação."),
        ("UC03 — Atualizar status", "Administrador", "Estar autenticado no Django Admin", "Localiza pedido, seleciona status permitido e salva.", "Novo status persistido com data de atualização.", "Valor fora das escolhas é rejeitado pelo formulário/modelo."),
    ]
    add_table(doc, ["Caso", "Ator", "Pré-condição", "Fluxo principal", "Pós-condição", "Alternativas"], expanded, [2.5, 2, 3, 5.2, 3.2, 3.6], 7)
    add_picture(doc, DIAGRAMS / "classes.png", caption="Figura 3 — Diagrama de classes do domínio.")
    add_picture(doc, DIAGRAMS / "sequencia-checkout.png", caption="Figura 4 — Diagrama de sequência do checkout.")
    add_picture(doc, DIAGRAMS / "atividades.png", caption="Figura 5 — Diagrama de atividades da jornada de compra.")

    add_heading(doc, "7. Arquitetura e tecnologias")
    add_body(doc, "A arquitetura é um monólito web com renderização no servidor. O navegador envia requisições HTTP às views Django; formulários validam entrada; o serviço de carrinho usa a sessão; o ORM persiste usuários, produtos, pedidos e itens. WhiteNoise serve arquivos estáticos na publicação, Gunicorn executa a aplicação e dj-database-url configura PostgreSQL.")
    add_table(doc, ["Camada", "Tecnologia", "Responsabilidade"], [
        ("Interface", "HTML5, CSS3, Bootstrap 5 e templates Django", "Navegação responsiva e acessível."),
        ("Aplicação", "Python 3.13 e Django 5.2 LTS", "Regras, autenticação, transações e administração."),
        ("Dados", "SQLite local / PostgreSQL no Render", "Persistência relacional."),
        ("Entrega", "GitHub Actions, Gunicorn, WhiteNoise e Render", "CI, execução e arquivos estáticos."),
    ], [3, 5.2, 9])

    add_heading(doc, "8. Interfaces, navegação e identidade visual")
    add_body(doc, "A marca combina rosa confeitaria, creme e chocolate, tipografia expressiva nos títulos e cartões claros. A hierarquia prioriza catálogo, benefícios e chamada para compra. A implementação evita dependências de imagens remotas: as fotografias demonstrativas foram geradas para o projeto e versionadas localmente.")
    add_picture(doc, DIAGRAMS / "wireframes.png", caption="Figura 6 — Wireframes das principais telas.")
    add_picture(doc, DIAGRAMS / "mapa-navegacional.png", caption="Figura 7 — Mapa navegacional.")
    add_picture(doc, SHOTS / "homepage-desktop.png", caption="Figura 8 — Homepage implementada em desktop (1440×1000).")
    add_picture(doc, SHOTS / "catalogo-desktop.png", caption="Figura 9 — Catálogo implementado em desktop.")
    add_picture(doc, SHOTS / "homepage-mobile.png", width=2.5, caption="Figura 10 — Homepage responsiva (390×844).")

    add_heading(doc, "9. Banco de dados")
    add_body(doc, "Modelo conceitual: um usuário realiza zero ou muitos pedidos; um pedido contém um ou muitos itens; cada item referencia opcionalmente um produto e preserva uma fotografia textual/financeira do momento da compra. Product e OrderItem não são fundidos porque representam catálogo mutável e histórico imutável, respectivamente.")
    add_picture(doc, DIAGRAMS / "modelo-logico.png", caption="Figura 11 — Modelo lógico normalizado.")
    add_heading(doc, "9.1 Modelo físico e dicionário", level=2)
    add_table(doc, ["Tabela.campo", "Tipo", "Restrições / significado"], [
        ("shop_product.id", "bigint", "PK"), ("shop_product.name", "varchar(120)", "Nome obrigatório"),
        ("shop_product.description", "text", "Descrição"), ("shop_product.price", "decimal(10,2)", ">= 0"),
        ("shop_product.image", "image/varchar", "Upload opcional"), ("shop_product.is_active", "boolean", "Disponibilidade"),
        ("shop_order.id", "bigint", "PK"), ("shop_order.user_id", "FK auth_user", "Proprietário"),
        ("shop_order.address/city/state/zip_code", "varchar", "Endereço de entrega"),
        ("shop_order.payment_method", "varchar(20)", "pix/card/cash"), ("shop_order.status", "varchar(24)", "Enum de operação"),
        ("shop_order.notes", "text", "Opcional"), ("shop_order.total", "decimal(10,2)", "Total congelado"),
        ("shop_order.created_at/updated_at", "datetime", "Auditoria"),
        ("shop_orderitem.order_id", "FK shop_order", "Pedido; cascade"), ("shop_orderitem.product_id", "FK shop_product", "Opcional; set null"),
        ("shop_orderitem.product_name", "varchar(120)", "Snapshot do nome"), ("shop_orderitem.unit_price", "decimal(10,2)", "Snapshot do preço"),
        ("shop_orderitem.quantity", "positive int", "1..20"), ("shop_orderitem.subtotal", "decimal(10,2)", "Preço × quantidade"),
    ], [5.2, 3.2, 9])

    add_heading(doc, "10. Testes e laudo de qualidade")
    add_body(doc, "A suíte automatizada executou 16 testes em modelos, formulários, autenticação, carrinho, totais, checkout transacional, permissões, página inicial, favicon e seed. Resultado final: 16 aprovados, zero falhas e 87% de cobertura total. O comando de verificação do Django não apontou problemas. O workflow do GitHub repete python manage.py test a cada push e pull request.")
    add_heading(doc, "10.1 Sessões internas de QA (simuladas)", level=2)
    add_body(doc, "As cinco sessões a seguir são avaliações internas simuladas para fins acadêmicos. Não representam colegas ou pessoas reais.")
    add_table(doc, ["Sessão", "Data", "Funcionou", "Achado e correção", "Não testado"], qa, [3, 2.2, 5, 5.5, 3], 7)
    add_heading(doc, "10.2 Laudo de correções", level=2)
    add_table(doc, ["Defeito", "Evidência", "Correção", "Verificação"], [
        ("Ordem visual dos destaques", "Teste inicialmente falhou porque o ORM ordenava alfabeticamente.", "Consulta da homepage passou a ordenar por criação.", "Teste específico aprovado."),
        ("Favicon ausente", "Console do navegador registrou 404.", "Favicon SVG local e link no template base.", "Console final: 0 erros e 0 avisos."),
        ("Risco de total manipulado", "Preço poderia ser adulterado no cliente sem revalidação.", "Carrinho e checkout recalculam pelo banco.", "Testes de total e checkout aprovados."),
        ("Exposição de pedido", "Risco de acesso por ID incremental.", "Consulta combina id e request.user; alheio retorna 404.", "Teste de isolamento aprovado."),
        ("Checkout parcial", "Falha intermediária poderia gerar pedido incompleto.", "Bloco transaction.atomic para pedido e itens.", "Teste de integração aprovado."),
    ], [3.2, 4.8, 5.3, 3.8], 7.5)

    add_heading(doc, "11. Instalação, execução e administração")
    add_heading(doc, "11.1 Execução local", level=2)
    add_bullets(doc, [
        "Instale Python 3.11+ e Git; clone o repositório.",
        "Crie ambiente virtual: python -m venv .venv; ative-o.",
        "Instale: pip install -r requirements.txt.",
        "Prepare: python manage.py migrate e python manage.py seed_demo.",
        "Opcional: python manage.py createsuperuser.",
        "Execute: python manage.py runserver; acesse http://127.0.0.1:8000/.",
        "Teste: python manage.py test; cobertura: coverage run manage.py test e coverage report.",
    ])
    add_heading(doc, "11.2 Manual do cliente", level=2)
    add_bullets(doc, [
        "Abra Produtos, escolha o cupcake e adicione-o ao carrinho.",
        "No carrinho, ajuste de 1 a 20 unidades ou remova itens.",
        "Entre/crie conta, abra Finalizar pedido e preencha endereço e pagamento demonstrativo.",
        "Confirme e acompanhe em Meus pedidos. O sistema não realiza cobrança real.",
    ])
    add_heading(doc, "11.3 Manual do administrador", level=2)
    add_bullets(doc, [
        "Acesse /admin/ com superusuário.",
        "Em Products, crie/edite itens, imagens, preços e disponibilidade.",
        "Em Orders, filtre pedidos e atualize apenas para um dos status permitidos.",
        "Não altere snapshots de itens históricos para refletir mudanças posteriores no catálogo.",
    ])
    add_heading(doc, "11.4 Publicação no Render", level=2)
    add_body(doc, "O repositório contém render.yaml e build.sh. No Render, conecte o GitHub, crie um Blueprint a partir do repositório, aceite o banco PostgreSQL e o serviço web e aguarde build/migrações. O ambiente recebe SECRET_KEY e DATABASE_URL automaticamente. Após a publicação, execute o seed demonstrativo pelo Shell ou ajuste o build conforme a política da disciplina. O plano gratuito pode hibernar e o banco gratuito possui limitações temporais; publique próximo à avaliação.")
    add_body(doc, f"Repositório público: {REPO_URL}\nURL pública: a inserir após a criação do Blueprint no Render.")

    add_heading(doc, "12. Roteiros dos vídeos")
    add_heading(doc, "12.1 Demonstração funcional — aproximadamente 5 minutos", level=2)
    add_table(doc, ["Tempo", "Narração / ação"], [
        ("0:00–0:30", "Apresentar o problema, a marca e o caráter acadêmico do pagamento."),
        ("0:30–1:15", "Mostrar homepage responsiva, benefícios e catálogo."),
        ("1:15–2:00", "Cadastrar usuário, entrar e explicar autenticação."),
        ("2:00–2:45", "Adicionar produtos, alterar quantidades, remover item e conferir total."),
        ("2:45–3:40", "Finalizar pedido com endereço e forma simulada; mostrar validação."),
        ("3:40–4:20", "Abrir histórico e detalhe; explicar isolamento entre clientes."),
        ("4:20–4:50", "Entrar no Admin, mostrar produto e atualização de status."),
        ("4:50–5:00", "Encerrar com GitHub, testes, cobertura e URL pública."),
    ], [3, 14.5])
    add_heading(doc, "12.2 Correções e qualidade — até 5 minutos", level=2)
    add_table(doc, ["Tempo", "Narração / ação"], [
        ("0:00–0:35", "Explicar que QA-01 a QA-05 são sessões internas simuladas."),
        ("0:35–1:25", "Apresentar achado de ordem dos destaques e teste que guiou a correção."),
        ("1:25–2:05", "Mostrar correção do favicon e console sem erros."),
        ("2:05–2:55", "Demonstrar limite 1–20 e recálculo de preços no servidor."),
        ("2:55–3:40", "Explicar transação do checkout e preservação do snapshot."),
        ("3:40–4:20", "Mostrar teste de acesso a pedido alheio retornando 404."),
        ("4:20–5:00", "Exibir 16 testes, 87% de cobertura e pipeline do GitHub."),
    ], [3, 14.5])

    add_heading(doc, "13. Conclusão")
    add_body(doc, "O Doce Encanto atende ao fluxo acadêmico definido: descoberta, seleção, checkout demonstrativo, histórico e operação administrativa. A solução mantém limites explícitos de escopo, cobre riscos centrais com testes e dispõe de documentação e configuração reproduzíveis. A publicação no Render requer apenas a vinculação da conta do aluno ao repositório público.")
    add_heading(doc, "Referências")
    add_bullets(doc, [
        "Django Software Foundation. Django 5.2 documentation. https://docs.djangoproject.com/en/5.2/",
        "Render. Deploy a Django App. https://render.com/docs/deploy-django",
        "Render. Free instances. https://render.com/docs/free",
        "Materiais e modelos do Projeto Integrador Transdisciplinar em Engenharia de Software I e II fornecidos pela instituição.",
    ])
    doc.save(path)


def build_pit_docx(path: Path) -> None:
    shutil.copy2(TEMPLATE, path)
    doc = Document(path)
    for sec in doc.sections:
        sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2); sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)
    replacements = {
        0: "Olá, estudante.",
        1: "A seguir, apresento a continuidade do projeto Doce Encanto e as respostas aos três desafios propostos.",
        3: "Nome: Matheus da Silva Mattos",
        4: "RGM: 36338770",
        6: "Documentação:",
        7: "A documentação da PIT I foi revisada e ampliada com histórias de usuário, backlog, requisitos, mapas, UML, modelos de dados, protótipos, telas finais, testes, manual, laudo e roteiros. O documento completo está versionado no repositório.",
        8: f"Link do repositório: {REPO_URL}",
        12: "Codificação:",
        13: "A tabela apresenta as tecnologias e os artefatos públicos da solução.",
        16: "Testes da Solução",
        17: "As cinco avaliações abaixo são sessões internas simuladas (QA-01 a QA-05), usadas somente para validação acadêmica. Não correspondem a colegas ou pessoas reais.",
        26: "Laudo de Qualidade",
        27: "Resultado final: 16 testes automatizados aprovados, zero falhas, cobertura total de 87% e console do navegador sem erros. Achados e correções estão resumidos abaixo.",
        28: "1. A ordem dos produtos em destaque não seguia a curadoria. Correção: ordenação por data de criação; teste específico aprovado.",
        29: "2. O favicon gerava resposta 404. Correção: inclusão de favicon SVG local; novo smoke test registrou zero erros e zero avisos.",
        30: "3. Os totais e limites precisam ser confiáveis. Correção: recálculo no servidor, quantidade entre 1 e 20 e bloqueio de produto inativo.",
        31: "4. Um cliente não pode abrir pedidos de outro. Correção: consulta filtrada simultaneamente por ID e usuário; acesso indevido retorna 404.",
        32: "5. Pedido e itens não podem ficar incompletos. Correção: checkout encapsulado em transação atômica, preservando nome e preço do produto.",
        33: "Vídeo da Solução atualizada",
        34: "O roteiro de até cinco minutos está incluído na documentação completa. O link abaixo deverá ser preenchido após a gravação e publicação pelo aluno.",
        35: "Link: a inserir após a gravação do vídeo.",
    }
    for idx, value in replacements.items():
        doc.paragraphs[idx].text = value
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.name = "Arial"; r.font.size = Pt(10)
    t = doc.tables[0]
    values = [
        "Python 3 / Django 5.2 LTS",
        "SQLite (local) e PostgreSQL (publicação)",
        "Render com Gunicorn e WhiteNoise",
        "Aplicação web responsiva",
        "(X) Tradicional\n( ) Low-code\n( ) No-code",
        REPO_URL,
        "A inserir após vincular o Blueprint no Render",
        "A inserir após gravar e publicar o vídeo narrado",
    ]
    for i, value in enumerate(values):
        set_cell_text(t.cell(i, 0), t.cell(i, 0).text, bold=True, size=8)
        set_cell_text(t.cell(i, 1), value, size=8)
        if i % 2 == 0:
            shade(t.cell(i, 0), CREAM); shade(t.cell(i, 1), CREAM)
    for qi, item in enumerate(qa, start=1):
        table = doc.tables[qi]
        name, dt, worked, issue, not_tested = item
        set_cell_text(table.cell(0, 0), f"Nome: {name}", bold=True, size=7.5)
        set_cell_text(table.cell(0, 1), f"Data do teste: {dt}", bold=True, size=7.5)
        set_cell_text(table.cell(1, 0), "O que testou e funcionou:", bold=True, size=7.5)
        set_cell_text(table.cell(1, 1), worked, size=7.5)
        set_cell_text(table.cell(2, 0), "Falha encontrada e correção:", bold=True, size=7.5)
        set_cell_text(table.cell(2, 1), issue, size=7.5)
        set_cell_text(table.cell(3, 0), "Funcionalidade não testada:", bold=True, size=7.5)
        set_cell_text(table.cell(3, 1), not_tested, size=7.5)
        for row in table.rows:
            shade(row.cells[0], CREAM)
    set_cell_text(doc.tables[6].cell(0, 0), "Link para o vídeo", bold=True, size=8)
    set_cell_text(doc.tables[6].cell(0, 1), "A inserir após a gravação e publicação pelo aluno", size=8)
    doc.save(path)


class NumberedDoc(BaseDocTemplate):
    def __init__(self, filename, **kw):
        super().__init__(filename, **kw)
        frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="normal")
        self.addPageTemplates(PageTemplate(id="all", frames=frame, onPage=self._page))

    def _page(self, canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(ROSE); canvas.line(2 * cm, A4[1] - 1.45 * cm, A4[0] - 2 * cm, A4[1] - 1.45 * cm)
        canvas.setFont("Helvetica", 7.5); canvas.setFillColor(COCOA)
        canvas.drawString(2 * cm, 1.15 * cm, "Doce Encanto • Matheus da Silva Mattos • RGM 36338770")
        canvas.drawRightString(A4[0] - 2 * cm, 1.15 * cm, f"Página {doc.page}")
        canvas.restoreState()


def pdf_styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("TitleX", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=26, leading=30, textColor=ROSE, alignment=TA_CENTER, spaceAfter=18),
        "sub": ParagraphStyle("SubX", parent=ss["Normal"], fontSize=13, leading=18, textColor=COCOA, alignment=TA_CENTER),
        "h1": ParagraphStyle("H1X", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=ROSE, spaceBefore=10, spaceAfter=7),
        "h2": ParagraphStyle("H2X", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=11.5, leading=15, textColor=COCOA, spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("BodyX", parent=ss["BodyText"], fontName="Helvetica", fontSize=8.8, leading=12, textColor=COCOA, alignment=TA_JUSTIFY, spaceAfter=5),
        "small": ParagraphStyle("SmallX", parent=ss["BodyText"], fontName="Helvetica", fontSize=7, leading=9, textColor=COCOA),
        "caption": ParagraphStyle("CaptionX", parent=ss["BodyText"], fontName="Helvetica-Oblique", fontSize=7, leading=9, textColor=COCOA, alignment=TA_CENTER, spaceAfter=7),
    }


def pdf_table(headers, rows, widths, styles, font=6.7):
    data = [[Paragraph(f"<b>{h}</b>", styles["small"]) for h in headers]]
    for row in rows:
        data.append([Paragraph(str(v).replace("\n", "<br/>"), styles["small"]) for v in row])
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="CENTER")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), ROSE), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#D8B7C2")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PALE]),
    ]))
    return table


def pdf_image(path: Path, styles, width=16.2 * cm, height=None, caption=""):
    if height is None:
        from PIL import Image as PILImage
        with PILImage.open(path) as im:
            height = width * im.height / im.width
        height = min(height, 19 * cm)
    return [Image(str(path), width=width, height=height), Paragraph(caption, styles["caption"])]


def build_main_pdf(path: Path) -> None:
    s = pdf_styles(); story = []
    story += [Spacer(1, 5.2 * cm), Paragraph("DOCE ENCANTO", s["title"]), Paragraph("E-commerce acadêmico de cupcakes<br/>Documentação de projeto e implementação", s["sub"]), Spacer(1, 2 * cm), Paragraph("Matheus da Silva Mattos<br/>RGM 36338770<br/><br/>São Paulo • 2026", s["sub"]), PageBreak()]
    def h(text): story.append(Paragraph(text, s["h1"]))
    def h2(text): story.append(Paragraph(text, s["h2"]))
    def body(text): story.append(Paragraph(text, s["body"]))
    h("Resumo executivo")
    body("O Doce Encanto é uma aplicação web responsiva para venda demonstrativa de cupcakes. Reúne cadastro, autenticação, catálogo, carrinho em sessão, checkout com pagamento simulado, histórico e administração. Foi desenvolvida em Django 5.2 LTS, com SQLite local e configuração PostgreSQL/Render.")
    body("Este relatório consolida análise, requisitos, backlog, UML, banco de dados, interface, testes, laudo, implantação, manuais e roteiros. O projeto é acadêmico e não processa pagamentos reais.")
    h("1. Visão, escopo e metodologia")
    body("Objetivo: demonstrar um fluxo de compra claro e seguro para uma confeitaria artesanal. Inclui contas, catálogo, carrinho, pedidos e administração; exclui gateway real, cupons, avaliações públicas, frete por distância, estoque complexo e aplicativo nativo.")
    story.append(pdf_table(["Etapa", "Atividade", "h"], tasks, [3.2*cm, 11.2*cm, 1.2*cm], s))
    story += pdf_image(DIAGRAMS / "mapa-afinidade.png", s, caption="Figura 1 — Mapa de afinidade.")
    h("2. Histórias de usuário")
    body("Prioridades e pontos orientaram a implementação. Regras transversais: preço recalculado no servidor, 1–20 unidades, produtos ativos, isolamento por usuário, transação atômica e pagamento demonstrativo.")
    story.append(pdf_table(["ID", "História", "Formulação e aceitação", "P", "Pts"], [(a,b,f"Como {c}, quero {d}, para {e}. {f}",g,hv) for a,b,c,d,e,f,g,hv in stories], [1.2*cm, 3*cm, 9.5*cm, 1.4*cm, 1.1*cm], s))
    h("3. Requisitos e regras de negócio")
    body("RF01–RF15 cobrem contas, catálogo, carrinho, checkout, pedidos, administração e tratamento de erros. RNF: responsividade, foco visível, variáveis de ambiente, PostgreSQL em produção, isolamento, reprodutibilidade e CI.")
    story.append(pdf_table(["Regra", "Descrição"], [
        ("RN01", "Quantidade por item entre 1 e 20."), ("RN02", "Somente produtos ativos podem ser comprados."),
        ("RN03", "Preços e totais são recalculados no servidor."), ("RN04", "Itens preservam nome e preço da compra."),
        ("RN05", "Checkout cria pedido e itens atomicamente."), ("RN06", "Pix simulado, cartão na entrega ou dinheiro."),
        ("RN07", "Status limitado ao fluxo operacional definido."), ("RN08", "Pedido só é visível ao proprietário/admin."),
    ], [2*cm, 14.8*cm], s))
    h("4. UML e fluxos")
    story += pdf_image(DIAGRAMS / "casos-de-uso.png", s, caption="Figura 2 — Casos de uso.")
    story.append(PageBreak())
    h2("Casos de uso expandidos")
    story.append(pdf_table(["Caso", "Ator", "Pré-condição", "Fluxo / resultado", "Alternativas"], [
        ("Cadastrar cliente", "Visitante", "Não autenticado", "Valida dados, cria usuário e inicia sessão.", "Duplicidade/senha inválida mantém o formulário."),
        ("Finalizar pedido", "Cliente", "Carrinho válido", "Revalida, cria pedido/itens e esvazia sessão.", "Vazio, inativo ou inválido não cria pedido."),
        ("Atualizar status", "Admin", "Acesso ao Admin", "Escolhe status permitido e salva.", "Status inválido é rejeitado."),
    ], [3*cm, 2*cm, 3*cm, 5*cm, 4*cm], s))
    story += pdf_image(DIAGRAMS / "classes.png", s, caption="Figura 3 — Classes.")
    story += pdf_image(DIAGRAMS / "sequencia-checkout.png", s, caption="Figura 4 — Sequência do checkout.")
    story += pdf_image(DIAGRAMS / "atividades.png", s, caption="Figura 5 — Atividades.")
    h("5. Arquitetura e interface")
    body("Monólito com templates renderizados no servidor: navegador → views/formulários/serviço de carrinho → ORM → banco. Gunicorn serve a aplicação e WhiteNoise os estáticos. A marca combina rosa, creme e chocolate e usa imagens locais demonstrativas.")
    story += pdf_image(DIAGRAMS / "wireframes.png", s, caption="Figura 6 — Wireframes.")
    story += pdf_image(DIAGRAMS / "mapa-navegacional.png", s, caption="Figura 7 — Navegação.")
    story += pdf_image(SHOTS / "homepage-desktop.png", s, caption="Figura 8 — Homepage desktop 1440×1000.")
    story += pdf_image(SHOTS / "catalogo-desktop.png", s, caption="Figura 9 — Catálogo desktop.")
    story += pdf_image(SHOTS / "homepage-mobile.png", s, width=6*cm, caption="Figura 10 — Homepage mobile 390×844.")
    h("6. Banco de dados")
    body("Um usuário possui muitos pedidos; cada pedido contém itens; o item referencia opcionalmente o catálogo e guarda snapshots. Essa separação mantém histórico mesmo se o produto mudar ou for removido.")
    story += pdf_image(DIAGRAMS / "modelo-logico.png", s, caption="Figura 11 — Modelo lógico.")
    story.append(pdf_table(["Entidade", "Campos centrais", "Integridade"], [
        ("Product", "name, description, price, image, is_active, timestamps", "price ≥ 0"),
        ("Order", "user, address, payment_method, status, total, timestamps", "FK usuário; escolhas controladas"),
        ("OrderItem", "order, product?, product_name, unit_price, quantity, subtotal", "1..20; snapshot histórico"),
        ("User", "username, email, password hash", "Autenticação nativa Django"),
    ], [3*cm, 9.5*cm, 4.5*cm], s))
    h("7. Testes e laudo de qualidade")
    body("Resultado final: 16 testes automatizados aprovados, zero falhas e 87% de cobertura. O Django check não apontou problemas. O workflow do GitHub executa a suíte em push e pull request.")
    story.append(pdf_table(["Sessão", "Funcionou", "Achado/correção", "Fora do escopo"], [(a,c,d,e) for a,b,c,d,e in qa], [3.2*cm, 4.6*cm, 6.1*cm, 3*cm], s))
    h2("Laudo")
    story.append(pdf_table(["Defeito", "Correção", "Verificação"], [
        ("Ordem dos destaques", "Ordenação por criação.", "Teste específico aprovado."),
        ("Favicon 404", "SVG local.", "Console: 0 erros/avisos."),
        ("Total manipulável", "Recálculo no servidor.", "Testes de total/checkout."),
        ("Pedido alheio", "Filtro por id + usuário.", "Retorno 404 testado."),
        ("Pedido parcial", "transaction.atomic.", "Integração aprovada."),
    ], [5*cm, 6*cm, 5.8*cm], s))
    h("8. Manual e publicação")
    body("Local: clone o repositório; crie e ative .venv; instale requirements.txt; execute migrate, seed_demo e runserver. Teste com python manage.py test. Cliente: selecione itens, ajuste o carrinho, entre, finalize com pagamento simulado e acompanhe em Meus pedidos. Admin: acesse /admin/, gerencie produtos e atualize status.")
    body(f"Repositório público: {REPO_URL}. Publicação: conecte o repositório no Render e crie o Blueprint definido em render.yaml. A URL será inserida após a autenticação do aluno no Render.")
    h("9. Roteiros")
    h2("Vídeo funcional — cerca de 5 minutos")
    story.append(pdf_table(["Tempo", "Ação"], [("0:00–0:30","Problema, marca e pagamento acadêmico."),("0:30–1:15","Homepage e catálogo."),("1:15–2:00","Cadastro/login."),("2:00–2:45","Carrinho e total."),("2:45–3:40","Checkout."),("3:40–4:20","Histórico e isolamento."),("4:20–4:50","Admin e status."),("4:50–5:00","GitHub, testes e URL.")], [3*cm, 13.8*cm], s))
    h2("Vídeo de correções — até 5 minutos")
    story.append(pdf_table(["Tempo", "Ação"], [("0:00–0:35","Explicar QA simulada."),("0:35–1:25","Ordem dos destaques."),("1:25–2:05","Favicon e console."),("2:05–2:55","Limites e totais."),("2:55–3:40","Transação e snapshots."),("3:40–4:20","Isolamento 404."),("4:20–5:00","16 testes, 87% e CI.")], [3*cm, 13.8*cm], s))
    h("10. Conclusão e referências")
    body("A solução atende ao fluxo planejado e cobre os principais riscos com validação no servidor, transações, permissões, testes e documentação. Referências: documentação Django 5.2; Render Deploy Django e Free Instances; materiais institucionais PIT I e PIT II.")
    doc = NumberedDoc(str(path), pagesize=A4, rightMargin=2*cm, leftMargin=2*cm, topMargin=1.8*cm, bottomMargin=1.7*cm, title="Doce Encanto — Documentação")
    doc.build(story)


def build_pit_pdf(path: Path) -> None:
    s = pdf_styles(); story = [Paragraph("PIT II — ATIVIDADE PREENCHIDA", s["title"]), Paragraph("Doce Encanto • Matheus da Silva Mattos • RGM 36338770", s["sub"]), Spacer(1, .6*cm)]
    story += [Paragraph("Documentação", s["h1"]), Paragraph(f"A documentação da PIT I foi revisada e ampliada. Repositório: {REPO_URL}", s["body"])]
    story += [Paragraph("Codificação", s["h1"]), pdf_table(["Campo", "Resposta"], [
        ("Back-end", "Python 3 / Django 5.2 LTS"), ("Banco", "SQLite local / PostgreSQL em produção"),
        ("Hospedagem", "Render + Gunicorn + WhiteNoise"), ("Plataforma", "Web responsiva"),
        ("Modo", "Tradicional"), ("GitHub", REPO_URL), ("Solução", "A inserir após vincular o Render"),
        ("Vídeo narrado", "A inserir após gravação pelo aluno"),
    ], [5*cm, 11.8*cm], s)]
    story += [Paragraph("Testes da solução", s["h1"]), Paragraph("QA-01 a QA-05 são sessões internas simuladas; não representam pessoas reais.", s["body"]), pdf_table(["Sessão/data", "Funcionou", "Falha/correção", "Não testado"], [(f"{a}<br/>{b}",c,d,e) for a,b,c,d,e in qa], [3.4*cm, 4.6*cm, 6*cm, 2.8*cm], s)]
    story += [Paragraph("Laudo de qualidade", s["h1"]), Paragraph("16 testes aprovados, zero falhas, 87% de cobertura e console final sem erros ou avisos. Correções: ordem de destaques por criação; favicon local; recálculo de preços e limite 1–20; isolamento por usuário; transação atômica no checkout.", s["body"])]
    story += [Paragraph("Vídeo da solução atualizada", s["h1"]), Paragraph("Link a inserir após a gravação e publicação pelo aluno. O roteiro de até cinco minutos está na documentação completa.", s["body"])]
    doc = NumberedDoc(str(path), pagesize=A4, rightMargin=2*cm, leftMargin=2*cm, topMargin=1.8*cm, bottomMargin=1.7*cm, title="PIT II preenchida")
    doc.build(story)


def main():
    OUTPUTS.mkdir(parents=True, exist_ok=True); DELIVERY.mkdir(parents=True, exist_ok=True)
    files = {
        "main_docx": OUTPUTS / "Documentacao_Projeto_Doce_Encanto.docx",
        "main_pdf": OUTPUTS / "Documentacao_Projeto_Doce_Encanto.pdf",
        "pit_docx": OUTPUTS / "PIT_atividade_preenchida.docx",
        "pit_pdf": OUTPUTS / "PIT_atividade_preenchida.pdf",
    }
    build_main_docx(files["main_docx"]); build_main_pdf(files["main_pdf"])
    build_pit_docx(files["pit_docx"]); build_pit_pdf(files["pit_pdf"])
    for src in files.values():
        shutil.copy2(src, DELIVERY / src.name)
    print("\n".join(str(p) for p in files.values()))


if __name__ == "__main__":
    main()
